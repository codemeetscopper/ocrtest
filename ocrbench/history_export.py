"""Tabular run-history downloads; never include frame images or raw records."""
import csv
import io
import json
from .terms import glossary


FIELDS = [
    'run_id', 'created_at', 'run_status', 'candidate', 'engine', 'requested_device',
    'observed_device', 'tesseract_binding', 'status', 'p95_ms', 'completed_fps',
    'e2e_f1', 'cer', 'peak_ram_mb', 'failures', 'completed', 'attempts', 'unique_samples',
    'median_ms', 'p99_ms', 'quality_scored_attempts', 'samples', 'mode', 'preprocessing',
    'track', 'languages', 'threads', 'repetitions', 'warmups', 'duration_s',
    'replay_fps', 'queue_policy', 'save_camera_frames', 'acceptance', 'engine_options', 'error',
]


def candidate_rows(run):
    config = run.get('config') or {}
    candidates = run.get('candidates') or [
        {'config': engine, 'status': run.get('status')}
        for engine in config.get('engines', [])
    ]
    rows = []
    for index, candidate in enumerate(candidates):
        engine = candidate.get('config') or {}
        summary = candidate.get('summary') or {}
        latency = summary.get('latency_ms') or {}
        identity = candidate.get('identity') or {}
        rows.append({
            'run_id': run['id'], 'created_at': run.get('created_at'), 'run_status': run.get('status'),
            'candidate': candidate.get('name') or f"{index}-{engine.get('engine', '')}-{engine.get('device', '')}",
            'engine': engine.get('engine'), 'requested_device': engine.get('device'),
            'observed_device': candidate.get('observed_device') or identity.get('observed_device'),
            'tesseract_binding': candidate.get('wrapper') or identity.get('wrapper') or (identity.get('runtime') or {}).get('binding'),
            'status': candidate.get('status'), 'p95_ms': latency.get('p95'),
            'completed_fps': summary.get('completed_fps'), 'e2e_f1': (summary.get('e2e') or {}).get('f1'),
            'cer': summary.get('cer'), 'peak_ram_mb': (candidate.get('resources') or {}).get('peak_process_tree_rss_mb'),
            **{key: summary.get(key) for key in ('failures', 'completed', 'attempts', 'unique_samples', 'quality_scored_attempts')},
            'median_ms': latency.get('median'), 'p99_ms': latency.get('p99'),
            'samples': ', '.join(config.get('samples') or []), 'languages': ', '.join(engine.get('languages') or []),
            **{key: config.get(key) for key in ('mode', 'preprocessing', 'track', 'repetitions', 'warmups',
                                              'duration_s', 'replay_fps', 'queue_policy', 'save_camera_frames')},
            'threads': engine.get('threads'), 'acceptance': json.dumps(summary.get('acceptance') or {}, ensure_ascii=False),
            'engine_options': json.dumps(engine.get('options') or {}, ensure_ascii=False), 'error': candidate.get('error'),
        })
    return rows


def csv_bytes(rows):
    out = io.StringIO(newline='')
    writer = csv.DictWriter(out, fieldnames=FIELDS)
    writer.writeheader()
    for row in rows:
        # Guard formula-like text even when preceded by whitespace. Numbers stay numeric.
        writer.writerow({key: "'"+value if isinstance(value, str) and
                         (value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n')))
                         else value for key, value in row.items()})
    return out.getvalue().encode('utf-8-sig')


def xlsx_bytes(rows):
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Run history'
    sheet.append(FIELDS)
    for row in rows:
        sheet.append([None]*len(FIELDS))
        for index, field in enumerate(FIELDS, 1):
            value = row.get(field)
            cell = sheet.cell(sheet.max_row, index)
            if isinstance(value, str):
                value = ILLEGAL_CHARACTERS_RE.sub('', value)
                if len(value) > 32767:
                    value = value[:32730]+'… [truncated for Excel]'
                cell.value = value
                cell.data_type = 's'  # Literal text, including strings starting with '='.
            else:
                cell.value = value
            if field in ('e2e_f1', 'cer'):
                cell.number_format = '0.00%'
            elif field in ('p95_ms', 'median_ms', 'p99_ms', 'completed_fps', 'peak_ram_mb'):
                cell.number_format = '0.00'
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30
    for index, field in enumerate(FIELDS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 28 if field in ('created_at', 'candidate', 'error', 'engine_options', 'samples') else max(15, len(field)+2)
    definitions = workbook.create_sheet('Metric definitions')
    definitions.append(['Term', 'Meaning'])
    definitions.append(['Export values', 'One row per engine result. Blank measurements are unavailable, never zero. CSV exports store F1 and CER as numeric ratios; Excel displays them as percentages. CER can exceed 100%. Timestamps retain the stored timezone.'])
    definitions.append(['Long text', 'Excel limits a cell to 32,767 characters. Longer text is marked as truncated; CSV retains the full text. Unsupported XML control characters are removed in Excel.'])
    for term in glossary():
        definitions.append([term['title'], term['description']])
    definitions.column_dimensions['A'].width = 35
    definitions.column_dimensions['B'].width = 110
    definitions.freeze_panes = 'B2'
    for row in definitions.iter_rows(min_row=2):
        definitions.row_dimensions[row[0].row].height = 60
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    for page in workbook:
        for cell in page[1]:
            cell.fill = PatternFill('solid', fgColor='244833')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(vertical='center')
        page.sheet_properties.pageSetUpPr.fitToPage = True
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
