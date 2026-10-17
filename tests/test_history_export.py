import csv
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from ocrbench.history_export import FIELDS
from ocrbench.schema import RunConfig
from ocrbench.storage import save_json


@pytest.fixture
def history_client(tmp_path, monkeypatch):
    import ocrbench.server as server
    monkeypatch.setattr(server, 'DATA', tmp_path)
    config = RunConfig(engines=[{'engine':'rapidocr'}, {'engine':'paddleocr','device':'gpu'}]).model_dump()
    candidate = {'name':'0-rapidocr-cpu','config':config['engines'][0],'status':'completed',
                 'identity':{'observed_device':'cpu'}, 'records':[{'predictions':[{'text':'Do not export frames'}]}],
                 'summary':{'latency_ms':{'p95':12.5,'p99':19.0,'median':10.0},'completed_fps':80.0,
                            'cer':0.0,'e2e':{'f1':0.9},'failures':0,'completed':5,'attempts':5},
                 'resources':{'peak_process_tree_rss_mb':123.4}}
    failed = {'name':'1-paddleocr-gpu','config':config['engines'][1],'status':'unavailable',
              'summary':None,'resources':None,'error':' \t=HYPERLINK("https://example.invalid", "仪表")'}
    for run_id, date, candidates in [('a'*12,'2026-10-15T12:00:00+00:00',[candidate,failed]),
                                      ('b'*12,'2026-10-16T12:00:00+00:00',[])]:
        save_json(tmp_path/'runs'/run_id/'run.json', {'id':run_id,'created_at':date,
                  'status':'completed_with_errors' if candidates else 'queued','config':config,'candidates':candidates})
    return TestClient(server.app, headers={'X-OCRBench':'local'})


def selected(*pairs):
    return {'rows':[{'run_id':run_id,'candidate_index':index} for run_id,index in pairs]}


def csv_rows(response):
    assert response.status_code == 200, response.text
    assert response.content.startswith(b'\xef\xbb\xbf')
    return list(csv.DictReader(io.StringIO(response.content.decode('utf-8-sig'))))


def test_csv_preserves_requested_subset_order_metrics_and_error_text(history_client):
    response = history_client.post('/api/history/export/csv',json=selected(('a'*12,1),('a'*12,0)))
    rows = csv_rows(response)
    assert len(rows) == 2 and [row['engine'] for row in rows] == ['paddleocr','rapidocr']
    assert rows[0]['p95_ms'] == '' and rows[0]['failures'] == ''
    assert rows[0]['error'].startswith("' \t=HYPERLINK") and '仪表' in rows[0]['error']
    assert rows[1]['cer'] == '0.0' and rows[1]['failures'] == '0'
    assert rows[1]['p95_ms'] == '12.5' and rows[1]['e2e_f1'] == '0.9'
    assert 'Do not export frames' not in response.text
    assert 'attachment; filename="ocrbench-history-' in response.headers['content-disposition']
    assert response.headers['content-type'].startswith('text/csv')


def test_xlsx_roundtrip_preserves_types_blanks_and_formula_like_text(history_client):
    response = history_client.post('/api/history/export/xlsx',json=selected(('a'*12,0),('a'*12,1)))
    assert response.status_code == 200
    assert response.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ['Run history','Metric definitions']
    sheet = workbook['Run history']
    assert sheet.max_row == 3 and sheet.freeze_panes == 'D2' and sheet.auto_filter.ref
    def cell(row, field):return sheet.cell(row,FIELDS.index(field)+1)
    assert cell(2,'cer').value == 0 and cell(2,'failures').value == 0
    assert cell(2,'e2e_f1').value == 0.9 and cell(2,'e2e_f1').number_format == '0.00%'
    assert cell(3,'p95_ms').value is None
    assert cell(3,'error').data_type == 's' and '仪表' in cell(3,'error').value
    assert not any(c.data_type == 'f' for row in sheet for c in row)
    assert any('P95 latency' in str(row[0].value) for row in workbook['Metric definitions'])
    workbook.close()


def test_all_history_includes_pending_engines_with_blank_metrics(history_client):
    rows = csv_rows(history_client.post('/api/history/export/csv',json={}))
    assert len(rows) == 4
    assert [r['run_id'] for r in rows] == ['b'*12,'b'*12,'a'*12,'a'*12]
    assert rows[0]['status'] == 'queued' and rows[0]['completed_fps'] == ''


@pytest.mark.parametrize('format',['csv','xlsx'])
def test_explicit_empty_selection_never_exports_all_runs(history_client,format):
    response = history_client.post('/api/history/export/'+format,json={'rows':[]})
    assert response.status_code == 200
    if format == 'csv':assert csv_rows(response) == []
    else:
        workbook = load_workbook(io.BytesIO(response.content))
        assert workbook['Run history'].max_row == 1
        workbook.close()


@pytest.mark.parametrize('body,code',[
    (selected(('c'*12,0)),404), (selected(('a'*12,9)),404),
    ({'rows':[{'run_id':'../README','candidate_index':0}]},422),
    ({'rows':[{'run_id':'a'*12,'candidate_index':-1}]},422),
    ({'rowz':[]},422),
])
def test_invalid_export_selection_is_rejected(history_client,body,code):
    assert history_client.post('/api/history/export/csv',json=body).status_code == code


def test_export_requires_local_header_and_valid_format(history_client):
    assert history_client.post('/api/history/export/csv',json={},headers={'X-OCRBench':''}).status_code == 403
    assert history_client.post('/api/history/export/pdf',json={}).status_code == 404


def test_excel_escapes_direct_formula_strings_and_invalid_xml_controls():
    from ocrbench.history_export import xlsx_bytes
    row = dict.fromkeys(FIELDS)
    row.update(run_id='a'*12,error='=1+1\x00',engine_options='@SUM(1,2)')
    workbook = load_workbook(io.BytesIO(xlsx_bytes([row])))
    sheet = workbook.active
    cell = sheet.cell(2,FIELDS.index('error')+1)
    assert cell.value == '=1+1' and cell.data_type == 's'
    workbook.close()
